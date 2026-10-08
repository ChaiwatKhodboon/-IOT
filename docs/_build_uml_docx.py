"""Build the report-ready Word copy of the IoT UML documentation."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
SPEC = ROOT / "uml-use-case-specifications.md"
IMAGES = ROOT / "uml-images"
OUTPUT = ROOT / "UML_IoT_Equipment_Report.docx"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    edges = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tbl_pr.append(edges)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, twips in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(twips))
        element.set(qn("w:type"), "dxa")
        mar.append(element)
    tc_pr.append(mar)


def set_cell_text(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    parts = value.replace("<br>", "\n").split("\n")
    paragraph = cell.paragraphs[0]
    paragraph.style = "Normal"
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    for i, part in enumerate(parts):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part.strip())
        run.bold = bold
        run.font.name = "Tahoma"
        run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def add_image(doc: Document, filename: str, caption: str, description: str) -> None:
    heading = doc.add_heading(caption, level=2)
    heading.paragraph_format.keep_with_next = True
    para = doc.add_paragraph(description)
    para.paragraph_format.keep_with_next = True
    path = IMAGES / filename
    with Image.open(path) as im:
        w, h = im.size
    max_width, max_height = 6.75, 7.75
    width = min(max_width, max_height * w / h)
    pic_para = doc.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_para.add_run().add_picture(str(path), width=Inches(width))
    pic_para.paragraph_format.space_after = Pt(10)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def create() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Tahoma"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Title"].font.size = Pt(21)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(12)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    for old_border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(old_border)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "UML ระบบยืมคืนอุปกรณ์ IoT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Tahoma"
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("หน้า ")
    add_page_field(footer)

    title = doc.add_paragraph("แผนภาพ UML ระบบยืมคืนอุปกรณ์ IoT", style="Title")
    title_ppr = title._p.get_or_add_pPr()
    for old_border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(old_border)
    doc.add_paragraph(
        "เอกสารฉบับนี้จัดทำจากระบบในโปรเจกต์ C:\\IOT เพื่อใช้ประกอบรายงานการวิเคราะห์และออกแบบระบบ "
        "ประกอบด้วย Use Case Diagram พร้อมคำอธิบายรายกรณี Class Diagram Sequence Diagram "
        "Activity Diagram และ State Diagram โดยแสดงกระบวนการยืม การส่งคำขอคืน การตรวจรับ "
        "การคืนบางส่วน และการซ่อมตามการทำงานจริงของระบบ"
    )
    doc.add_paragraph(
        "ผู้ใช้ทั่วไปค้นหาอุปกรณ์และยืมได้ทันทีเมื่อมีจำนวนพร้อมยืม เมื่อแจ้งคืน ระบบจะรอให้ผู้ดูแลตรวจรับของจริง "
        "ก่อนปรับสต็อก ผู้ดูแลยังจัดการคลัง ผู้ใช้ รายงาน และประวัติการทำงานได้"
    )
    doc.add_heading("ขอบเขตของแผนภาพ", level=1)
    for text in (
        "Use Case Diagram แสดงผู้เกี่ยวข้องและงานที่ทำกับระบบ แบ่งภาพผู้ใช้กับผู้ดูแลเพื่อให้ข้อความอ่านได้",
        "Class Diagram แสดงข้อมูลและความสัมพันธ์หลัก รวมทั้งพฤติกรรมเชิงแนวคิดของคลาส",
        "Sequence Diagram แสดงทั้งการโต้ตอบระหว่าง Actor กับระบบและลำดับการทำงานภายใน",
        "Activity Diagram แสดงขั้นตอนยืม คืน ตรวจรับ และซ่อมตามผู้รับผิดชอบ",
        "State Diagram แยกวงจรรายการยืมกับส่วนคืนตามสภาพ",
    ):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph(
        "หมายเหตุ: ในระบบจริง pending_return เป็นข้อมูลประกอบรายการ borrowed ไม่ใช่ค่า status ในฐานข้อมูล "
        "และการเกินกำหนดคำนวณจากวันที่กำหนดคืน ไม่ใช่สถานะที่บันทึกแยก"
    )

    doc.add_page_break()
    doc.add_heading("แผนภาพ UML", level=1)
    diagrams = [
        ("uml-use-case-member.png", "ภาพที่ 1 Use Case สำหรับผู้เยี่ยมชมและผู้ใช้", "แสดงการสมัคร เข้าสู่ระบบ ดูคลัง ยืม ส่งคำขอคืน และติดตามรายการ"),
        ("uml-use-case-admin.png", "ภาพที่ 2 Use Case สำหรับผู้ดูแล", "แสดงงานเพิ่มเติมของผู้ดูแล ได้แก่ คลัง ตรวจรับ ซ่อม ผู้ใช้ รายงาน และ Audit"),
        ("uml-class.png", "ภาพที่ 3 Class Diagram", "แสดงคลาสข้อมูล ความสัมพันธ์ จำนวนคู่สัมพันธ์ และหน้าที่เชิงแนวคิด"),
        ("uml-sequence-system-return.png", "ภาพที่ 4 System Sequence Diagram การคืนและตรวจรับ", "แสดงคำสั่งและผลตอบกลับระหว่างผู้ยืม ผู้ดูแล และระบบในระดับ OOA"),
        ("uml-sequence-borrow.png", "ภาพที่ 5 Sequence Diagram การยืม", "แสดงการตรวจสิทธิ์ ล็อกสต็อก บันทึกยืม และกรณีจำนวนไม่พอ"),
        ("uml-sequence-return-inspection.png", "ภาพที่ 6 Sequence Diagram คำขอคืนและตรวจรับ", "แสดงการส่งคำขอ การส่งกลับแก้ไข และการตรวจรับตามสภาพของจริง"),
        ("uml-sequence-repair.png", "ภาพที่ 7 Sequence Diagram การปิดงานซ่อม", "แสดงการย้ายจำนวนจากรอซ่อมกลับไปพร้อมยืม"),
        ("uml-activity.png", "ภาพที่ 8 Activity Diagram", "แสดงลำดับกิจกรรมตั้งแต่ยืมจนคืนและซ่อม พร้อมผู้รับผิดชอบแต่ละขั้น"),
        ("uml-state-loan.png", "ภาพที่ 9 State Diagram รายการยืม", "แสดงกำลังยืม รอตรวจรับ คืนครบ และเหตุการณ์ที่เปลี่ยนสถานะ"),
        ("uml-state-return.png", "ภาพที่ 10 State Diagram ส่วนคืน", "แสดงส่วนคืนปกติ ชำรุดหรือผิดปกติ สูญหาย และซ่อมเสร็จ"),
    ]
    for file, caption, description in diagrams:
        add_image(doc, file, caption, description)

    doc.add_page_break()
    doc.add_heading("คำอธิบาย Use Case รายตัว", level=1)
    doc.add_paragraph(
        "ตารางต่อไปนี้ใช้รหัสเดียวกับ Use Case Diagram ทุกกรณี ระบุผู้กระทำ จุดเริ่ม เงื่อนไข "
        "Main Flow Exceptional Flow และผลหลังดำเนินการ"
    )
    source = SPEC.read_text(encoding="utf-8")
    matches = list(re.finditer(r"(?m)^## (UC\d{2} .+)$", source))
    assert len(matches) == 22
    for i, match in enumerate(matches):
        title = match.group(1).strip()
        body = source[match.end() : matches[i + 1].start() if i + 1 < len(matches) else len(source)]
        rows = [line.strip() for line in body.splitlines() if line.startswith("| ")]
        records = []
        for row in rows:
            cells = [part.strip() for part in row.strip("|").split("|")]
            if len(cells) == 2 and cells[0] not in ("หัวข้อ", "---"):
                records.append(cells)
        assert len(records) == 8, (title, records)
        heading = doc.add_heading(title, level=2)
        heading.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.45)
        table.columns[1].width = Inches(5.45)
        borders(table)
        header_row = table.rows[0]
        repeat_header(header_row)
        keep_row_together(header_row)
        set_cell_text(header_row.cells[0], "หัวข้อ", bold=True)
        set_cell_text(header_row.cells[1], "รายละเอียด", bold=True)
        for cell in header_row.cells:
            shade(cell, "EAF0F5")
        for idx, (key, value) in enumerate(records):
            row = table.add_row()
            keep_row_together(row)
            set_cell_text(row.cells[0], key, bold=True)
            set_cell_text(row.cells[1], value)
            if idx % 2:
                for cell in row.cells:
                    shade(cell, "F8FAFC")
        doc.add_paragraph("")

    doc.add_heading("ข้อมูลอ้างอิง", level=1)
    doc.add_paragraph(
        "แผนภาพและคำอธิบายอ้างอิง README.md, database/schema.sql, migrations 011 และ 012, "
        "src/routes และเอกสารกระบวนการตรวจรับในโปรเจกต์ ส่วนรูปแบบการเขียน UML "
        "ตรวจเทียบกับไฟล์บทเรียน OOA/OOD ที่ผู้ใช้ให้มา"
    )
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create()
